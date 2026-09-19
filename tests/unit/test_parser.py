from docx import Document

from src.ingestion.parser import docx_to_sections


def test_docx_to_sections_extracts_sections(tmp_path):
    path = tmp_path / "test.docx"

    document = Document()
    document.add_heading("Motor Mounting", level=1)
    document.add_paragraph("Mount the motors.")
    document.add_heading("Battery Selection", level=1)
    document.add_paragraph("Use a 4S battery.")
    document.save(path)

    sections = docx_to_sections(path)

    assert sections == [
        {
            "heading": "Motor Mounting",
            "content": ["Mount the motors."],
        },
        {
            "heading": "Battery Selection",
            "content": ["Use a 4S battery."],
        },
    ]


def test_docx_to_sections_groups_paragraphs_under_heading(tmp_path):
    path = tmp_path / "test.docx"

    document = Document()
    document.add_heading("Motor Mounting", level=1)
    document.add_paragraph("Mount the motors.")
    document.add_paragraph("Secure the screws.")
    document.save(path)

    sections = docx_to_sections(path)

    assert sections[0]["content"] == [
        "Mount the motors.",
        "Secure the screws.",
    ]


def test_docx_to_sections_ignores_empty_paragraphs(tmp_path):
    path = tmp_path / "test.docx"

    document = Document()
    document.add_heading("Motor Mounting", level=1)
    document.add_paragraph("")
    document.add_paragraph("       ")
    document.add_paragraph("Mount the motors.")
    document.add_paragraph("")
    document.save(path)

    sections = docx_to_sections(path)

    assert sections[0]["content"] == ["Mount the motors."]
